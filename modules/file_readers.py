from pathlib import Path

import pandas as pd
from pypdf import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {".csv", ".pdf", ".docx"}


def read_csv(file_path: str | Path) -> pd.DataFrame:
    """Read a CSV file as structured tabular data."""
    path = Path(file_path)

    if path.suffix.lower() != ".csv":
        raise ValueError("Expected a CSV file.")

    return pd.read_csv(path)


def read_pdf(file_path: str | Path) -> str:
    """Extract selectable text from a PDF file."""
    path = Path(file_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError("Expected a PDF file.")

    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            pages.append(f"[Page {page_number}]\n{text.strip()}")

    return "\n\n".join(pages)


def read_docx(file_path: str | Path) -> str:
    """Extract paragraphs and table text from a Word document."""
    path = Path(file_path)

    if path.suffix.lower() != ".docx":
        raise ValueError("Expected a DOCX file.")

    document = Document(str(path))
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        rows = []

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))

        if rows:
            sections.append(
                f"[Table {table_number}]\n" + "\n".join(rows)
            )

    return "\n\n".join(sections)


def read_file(file_path: str | Path):
    """Read a supported file using its extension."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            "Supported types: CSV, PDF, DOCX."
        )

    if extension == ".csv":
        return read_csv(path)

    if extension == ".pdf":
        return read_pdf(path)

    return read_docx(path)
